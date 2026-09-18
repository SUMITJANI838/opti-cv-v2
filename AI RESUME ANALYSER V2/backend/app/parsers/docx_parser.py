import io
from typing import Tuple, Dict, Any
from docx import Document

def extract_text_from_docx(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """
    Extracts text content and paragraph/table structure from DOCX byte stream.
    """
    try:
        doc = Document(io.BytesIO(file_bytes))
        text_chunks = []
        
        # Paragraphs
        for p in doc.paragraphs:
            if p.text.strip():
                text_chunks.append(p.text.strip())
                
        # Tables (e.g. skills matrix or experience in tables)
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_chunks.append(" | ".join(row_text))
                    
        full_text = "\n".join(text_chunks)
        metadata = {
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
        }
        return full_text, metadata
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX document: {str(e)}")
