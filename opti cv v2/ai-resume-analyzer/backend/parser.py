import re
import io
from typing import Dict, List, Any, Optional
from pypdf import PdfReader
from docx import Document

# Comprehensive Skills Taxonomy
SKILLS_TAXONOMY = {
    "languages": [
        "python", "javascript", "typescript", "c++", "c#", "java", "rust", "go", "golang",
        "ruby", "php", "swift", "kotlin", "scala", "r", "dart", "sql", "html", "css", "bash", "shell"
    ],
    "frontend": [
        "react", "react.js", "next.js", "nextjs", "vue", "vue.js", "nuxt", "angular",
        "svelte", "sveltekit", "tailwind", "tailwind css", "bootstrap", "sass", "redux",
        "zustand", "framer motion", "three.js", "graphql", "html5", "css3", "webpack", "vite"
    ],
    "backend": [
        "node.js", "nodejs", "express", "fastapi", "flask", "django", "spring boot", "nestjs",
        "ruby on rails", "asp.net", "laravel", "gin", "grpc", "rest api", "restful api", "microservices"
    ],
    "databases": [
        "postgresql", "postgres", "mysql", "mongodb", "redis", "elasticsearch", "sqlite",
        "dynamodb", "cassandra", "neo4j", "supabase", "firebase", "firestore", "mariadb", "oracle"
    ],
    "cloud_devops": [
        "aws", "amazon web services", "azure", "google cloud", "gcp", "docker", "kubernetes",
        "k8s", "terraform", "ci/cd", "github actions", "gitlab ci", "jenkins", "ansible",
        "linux", "nginx", "prometheus", "grafana", "helm", "serverless", "lambda"
    ],
    "ai_ml": [
        "machine learning", "deep learning", "nlp", "computer vision", "pytorch", "tensorflow",
        "keras", "scikit-learn", "huggingface", "llm", "large language models", "rag",
        "langchain", "llamaindex", "openai", "gemini", "pandas", "numpy", "opencv", "transformers"
    ],
    "soft_skills": [
        "leadership", "communication", "collaboration", "agile", "scrum", "problem solving",
        "critical thinking", "mentorship", "project management", "cross-functional", "adaptability"
    ]
}

SECTION_HEADERS = {
    "summary": ["summary", "professional summary", "about me", "profile", "objective", "executive summary"],
    "experience": ["experience", "work experience", "professional experience", "employment history", "work history"],
    "projects": ["projects", "personal projects", "key projects", "notable projects", "academic projects"],
    "skills": ["skills", "technical skills", "core competencies", "technologies", "tech stack", "tools & technologies"],
    "education": ["education", "academic background", "academic history", "qualifications", "degrees"],
    "certifications": ["certifications", "licenses", "certificates", "accreditations"],
    "awards": ["awards", "honors", "achievements", "accomplishments", "publications"]
}

ACTION_VERBS = [
    "architected", "developed", "engineered", "designed", "implemented", "spearheaded",
    "optimized", "accelerated", "scaled", "automated", "built", "reduced", "increased",
    "improved", "orchestrated", "deployed", "transformed", "streamlined", "enhanced",
    "delivered", "launched", "pioneered", "mentored", "led", "boosted", "integrated"
]

def extract_text_from_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    text_parts = []
    for page_idx, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        if page_text.strip():
            text_parts.append(page_text.strip())
    return "\n\n".join(text_parts)

def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    text_parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            text_parts.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                text_parts.append(" | ".join(row_text))
    return "\n".join(text_parts)

def extract_contact_info(text: str) -> Dict[str, Any]:
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    
    # Candidate Name (heuristic: first non-empty line without special characters/emails)
    name = "Candidate"
    for line in lines[:5]:
        if not re.search(r"[@:/\\_0-9|•]", line) and 2 <= len(line.split()) <= 4 and len(line) < 40:
            name = line.strip()
            break

    # Email
    email_match = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    email = email_match.group(0) if email_match else None

    # Phone
    phone_match = re.search(r"(\+?\d{1,3}[-.\s]?)?(\(?\d{3}\)?[-.\s]?)?\d{3}[-.\s]?\d{4}", text)
    phone = phone_match.group(0) if phone_match else None

    # Links
    linkedin_match = re.search(r"(linkedin\.com/in/[\w-]+)", text, re.IGNORECASE)
    linkedin = f"https://{linkedin_match.group(1)}" if linkedin_match else None

    github_match = re.search(r"(github\.com/[\w-]+)", text, re.IGNORECASE)
    github = f"https://{github_match.group(1)}" if github_match else None

    portfolio_match = re.search(r"(https?://[\w.-]+(?:portfolio|dev|tech|io|me|com)/?[\w/-]*)", text, re.IGNORECASE)
    portfolio = portfolio_match.group(1) if portfolio_match else None

    return {
        "name": name,
        "email": email,
        "phone": phone,
        "linkedin": linkedin,
        "github": github,
        "portfolio": portfolio
    }

def segment_sections(text: str) -> Dict[str, str]:
    lines = text.split("\n")
    sections: Dict[str, List[str]] = {
        "summary": [],
        "experience": [],
        "projects": [],
        "skills": [],
        "education": [],
        "certifications": [],
        "other": []
    }
    
    current_section = "summary"
    
    for line in lines:
        cleaned = line.strip()
        if not cleaned:
            continue
        
        # Check if this line is a section header
        header_candidate = re.sub(r"[^\w\s]", "", cleaned.lower()).strip()
        found_section = None
        
        for section_key, aliases in SECTION_HEADERS.items():
            if header_candidate in aliases or (len(header_candidate.split()) <= 3 and any(alias in header_candidate for alias in aliases)):
                found_section = section_key
                break
        
        if found_section:
            current_section = found_section
        else:
            if current_section in sections:
                sections[current_section].append(cleaned)
            else:
                sections["other"].append(cleaned)
                
    return {k: "\n".join(v).strip() for k, v in sections.items()}

def extract_skills(text: str) -> Dict[str, List[str]]:
    text_lower = text.lower()
    detected_skills: Dict[str, List[str]] = {}
    
    for category, skills_list in SKILLS_TAXONOMY.items():
        found = []
        for skill in skills_list:
            # Word boundary check
            pattern = r"(?<!\w)" + re.escape(skill) + r"(?!\w)"
            if re.search(pattern, text_lower):
                found.append(skill.title() if len(skill) > 3 else skill.upper())
        if found:
            detected_skills[category] = list(set(found))
            
    return detected_skills

def extract_bullets(text: str) -> List[str]:
    lines = text.split("\n")
    bullets = []
    bullet_pattern = r"^[\s•\-*>\d\.]+\s*(.+)$"
    
    for line in lines:
        cleaned = line.strip()
        if not cleaned:
            continue
        # Check if line looks like a bullet or description
        if any(cleaned.startswith(p) for p in ["•", "-", "*", "–", "—", ">"]) or re.match(r"^\d+[\.\)]\s+", cleaned):
            match = re.match(bullet_pattern, cleaned)
            if match and len(match.group(1)) > 15:
                bullets.append(match.group(1).strip())
        elif len(cleaned) > 35 and not any(cleaned.lower().startswith(h) for h in ["education", "skills", "experience", "projects", "certifications"]):
            # Longer standalone sentence in experience/projects
            bullets.append(cleaned)
            
    return bullets

def parse_resume_file(filename: str, content_bytes: bytes) -> Dict[str, Any]:
    ext = filename.split(".")[-1].lower()
    raw_text = ""
    
    if ext == "pdf":
        raw_text = extract_text_from_pdf(content_bytes)
    elif ext in ["docx", "doc"]:
        raw_text = extract_text_from_docx(content_bytes)
    elif ext in ["txt", "md"]:
        raw_text = content_bytes.decode("utf-8", errors="ignore")
    else:
        # Fallback decode
        raw_text = content_bytes.decode("utf-8", errors="ignore")

    raw_text = raw_text.strip()
    if not raw_text:
        raise ValueError("Could not extract any readable text from the file.")

    contact = extract_contact_info(raw_text)
    sections = segment_sections(raw_text)
    skills = extract_skills(raw_text)
    
    # Extract bullets from experience and projects
    combined_exp = sections.get("experience", "") + "\n" + sections.get("projects", "")
    bullets = extract_bullets(combined_exp if combined_exp.strip() else raw_text)
    
    word_count = len(raw_text.split())
    char_count = len(raw_text)
    
    return {
        "filename": filename,
        "word_count": word_count,
        "char_count": char_count,
        "contact": contact,
        "sections": sections,
        "skills": skills,
        "bullets": bullets,
        "raw_text": raw_text
    }
